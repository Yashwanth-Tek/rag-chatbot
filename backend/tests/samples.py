"""Builders for test inputs: real PDF/DOCX bytes and a deterministic stand-in for Ollama."""

import io
import math
import re
import zlib
from types import SimpleNamespace

from docx import Document

from app.errors import AppError


def make_pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF with one line of Helvetica text per page (empty string = blank page)."""
    objects: list[bytes] = []
    page_ids = [3 + 2 * i for i in range(len(pages))]
    font_id = 3 + 2 * len(pages)
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids).encode()
    objects.append(b"<< /Type /Pages /Kids [" + kids + b"] /Count %d >>" % len(pages))
    for index, text in enumerate(pages):
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 11 Tf 50 750 Td ({escaped}) Tj ET".encode("latin-1") if text else b""
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792]"
            b" /Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>"
            % (font_id, page_ids[index] + 1)
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    out.writelines(b"%010d 00000 n \n" % offset for offset in offsets)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    )
    return out.getvalue()


def make_docx(blocks: list[tuple[str, str]]) -> bytes:
    """blocks: ("h1".."h3", text) for headings, ("p", text) for paragraphs,
    ("table", "a|b;c|d") for a table with rows separated by ';' and cells by '|'."""
    document = Document()
    for kind, text in blocks:
        if kind.startswith("h"):
            document.add_heading(text, level=int(kind[1]))
        elif kind == "table":
            rows = [row.split("|") for row in text.split(";")]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, cells in enumerate(rows):
                for c, value in enumerate(cells):
                    table.cell(r, c).text = value
        else:
            document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class FakeEmbedder:
    """Bag-of-words hashing into a fixed dimension: texts sharing words get similar vectors,
    which is enough to exercise retrieval without running Ollama."""

    model_id = "fake/bag-of-words"

    def __init__(self, dimension: int = 768) -> None:
        self.dimension = dimension
        self.calls = 0

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        self.calls += 1
        return [self._vector(text) for text in texts]

    def embed_query(self, text: str) -> list[float]:
        self.calls += 1
        return self._vector(text)

    def check(self) -> str | None:
        return None

    def _vector(self, text: str) -> list[float]:
        vector = [0.0] * self.dimension
        for word in re.findall(r"[a-z0-9]+", text.lower()):
            vector[zlib.crc32(word.encode()) % self.dimension] += 1.0
        norm = math.sqrt(sum(v * v for v in vector)) or 1.0
        return [v / norm for v in vector]


AnswerSpec = list[tuple[str, int | None]] | str  # [(text, cited chunk index or None)] | "refusal"


class FakeClaude:
    """Stands in for anthropic.Anthropic with scripted responses.

    answers:  one spec per answer request (the last repeats), e.g.
              [("The company was founded in 2015.", 0), (" The CEO is unknown.", None)]
              A cited block quotes the whole chunk as its cited_text.
    verdicts: one verdict list per verification request (the last repeats).
    """

    def __init__(self, answers: list[AnswerSpec] = (), verdicts: list[list[str]] = ()) -> None:
        self.answers = list(answers)
        self.verdicts = list(verdicts)
        self.answer_requests: list[dict] = []
        self.verify_requests: list[dict] = []
        self.beta = SimpleNamespace(
            messages=SimpleNamespace(create=self._answer, parse=self._verify)
        )
        self.models = SimpleNamespace(retrieve=lambda model: SimpleNamespace(id=model))

    def with_options(self, **_options) -> "FakeClaude":
        return self

    def _answer(self, **request) -> SimpleNamespace:
        self.answer_requests.append(request)
        assert self.answers, "Claude was called but the test expected it not to be"
        spec = self.answers[min(len(self.answer_requests), len(self.answers)) - 1]
        if spec == "refusal":
            return SimpleNamespace(stop_reason="refusal", content=[])
        documents = [b for b in request["messages"][0]["content"] if b["type"] == "document"]
        blocks = [
            SimpleNamespace(
                type="text",
                text=text,
                citations=None
                if index is None
                else [
                    SimpleNamespace(
                        type="char_location",
                        document_index=index,
                        cited_text=documents[index]["source"]["data"],
                    )
                ],
            )
            for text, index in spec
        ]
        return SimpleNamespace(stop_reason="end_turn", content=blocks)

    def _verify(self, **request) -> SimpleNamespace:
        self.verify_requests.append(request)
        assert self.verdicts, "verifier was called but the test gave no verdicts"
        verdicts = self.verdicts[min(len(self.verify_requests), len(self.verdicts)) - 1]
        parsed = SimpleNamespace(
            verdicts=[
                SimpleNamespace(sentence=n, verdict=v, reason="")
                for n, v in enumerate(verdicts, start=1)
            ]
        )
        return SimpleNamespace(stop_reason="end_turn", parsed_output=parsed)


class UnavailableEmbedder(FakeEmbedder):
    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        raise AppError("EMBEDDING_UNAVAILABLE", "The embedding service is unavailable.", 503)

    def embed_query(self, text: str) -> list[float]:
        raise AppError("EMBEDDING_UNAVAILABLE", "The embedding service is unavailable.", 503)

    def check(self) -> str | None:
        return "Ollama isn't reachable."
